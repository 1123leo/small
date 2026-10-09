#!/usr/bin/env python3
"""Convert Markdown report to Word (.docx) format"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import re

def markdown_to_docx(md_file, docx_file):
    """Convert markdown file to Word document"""
    
    # Read markdown content
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create Document
    doc = Document()
    
    # Process markdown content line by line
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # Skip empty lines
        if not line.strip():
            i += 1
            continue
        
        # H1 - Title
        if line.startswith('# '):
            title = line[2:].strip()
            heading = doc.add_heading(title, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        
        # H2 - Main sections
        if line.startswith('## '):
            section = line[3:].strip()
            doc.add_heading(section, level=2)
            i += 1
            continue
        
        # H3 - Subsections
        if line.startswith('### '):
            subsection = line[4:].strip()
            doc.add_heading(subsection, level=3)
            i += 1
            continue
        
        # H4
        if line.startswith('#### '):
            subsubsection = line[5:].strip()
            doc.add_heading(subsubsection, level=4)
            i += 1
            continue
        
        # Tables
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            
            if table_lines:
                # Parse table
                rows = []
                for table_line in table_lines:
                    cells = [cell.strip() for cell in table_line.split('|')[1:-1]]
                    if cells and not all(c == '-' * len(c) or c.replace('-', '').replace(':', '') == '' for c in cells):
                        rows.append(cells)
                
                if rows:
                    # Create table
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = 'Light Grid Accent 1'
                    
                    # Fill table
                    for row_idx, row_data in enumerate(rows):
                        for col_idx, cell_data in enumerate(row_data):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_data
                            # Bold header row
                            if row_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.font.bold = True
            continue
        
        # Code blocks
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1  # Skip closing ```
            
            if code_lines:
                # Add code block
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph(code_text, style='No Spacing')
                p_format = p.paragraph_format
                p_format.left_indent = Inches(0.5)
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
            continue

        # Images
        image_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', line.strip())
        if image_match:
            alt_text = image_match.group(1).strip()
            image_path = Path(image_match.group(2).strip())
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(image_path), width=Inches(5.5))
                if alt_text:
                    caption = doc.add_paragraph(alt_text)
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in caption.runs:
                        run.font.size = Pt(10)
                        run.font.italic = True
            else:
                doc.add_paragraph(f"[圖片未找到：{image_path}]")
            i += 1
            continue
        
        # Bullet lists
        if line.strip().startswith('- '):
            bullet_lines = []
            while i < len(lines) and lines[i].strip().startswith('- '):
                bullet_lines.append(lines[i].strip()[2:])
                i += 1
            
            for bullet in bullet_lines:
                doc.add_paragraph(bullet, style='List Bullet')
            continue
        
        # Numbered lists
        if re.match(r'^\d+\.\s', line.strip()):
            num_lines = []
            while i < len(lines) and re.match(r'^\d+\.\s', lines[i].strip()):
                num_lines.append(lines[i].strip().split('. ', 1)[1])
                i += 1
            
            for num, text in enumerate(num_lines, 1):
                doc.add_paragraph(text, style='List Number')
            continue
        
        # Regular paragraphs
        if line.strip():
            paragraph = doc.add_paragraph(line.strip())
            # Try to format inline markdown
            if '**' in line or '*' in line or '`' in line:
                # Clear and reformat
                paragraph.clear()
                parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', line.strip())
                for part in parts:
                    if not part:
                        continue
                    run = paragraph.add_run(part)
                    if part.startswith('**') and part.endswith('**'):
                        run.text = part[2:-2]
                        run.font.bold = True
                    elif part.startswith('`') and part.endswith('`'):
                        run.text = part[1:-1]
                        run.font.name = 'Courier New'
                        run.font.color.rgb = RGBColor(192, 0, 0)
                    elif part.startswith('*') and part.endswith('*'):
                        run.text = part[1:-1]
                        run.font.italic = True
            i += 1
            continue
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"✅ Successfully converted to: {docx_file}")
    return True

# Convert report
md_path = Path('flappy_bird_ppo_report.md')
docx_path = Path('flappy_bird_ppo_report.docx')

if md_path.exists():
    print(f"Converting {md_path.name} to Word format...")
    markdown_to_docx(str(md_path), str(docx_path))
    
    # Verify
    if docx_path.exists():
        size_kb = docx_path.stat().st_size / 1024
        print(f"📄 Word 檔案大小：{size_kb:.1f} KB")
        print(f"📍 位置：{docx_path.absolute()}")
else:
    print(f"❌ {md_path} not found")
